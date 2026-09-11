import csv
import json
from io import StringIO, BytesIO
from typing import Optional
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_, func

from app.api.deps import get_db, verify_admin_key
from app.models.session import Session
from app.models.event import TelemetryEvent
from app.models.survey import SurveyResponse
from app.services.export import stream_csv, stream_json, stream_jsonl

router = APIRouter(dependencies=[Depends(verify_admin_key)])

def format_csv_row(row: list) -> str:
    si = StringIO()
    cw = csv.writer(si)
    cw.writerow(row)
    return si.getvalue()

async def stream_student_summary_csv(db: AsyncSession):
    headers = [
        "participant_id", "condition", "has_assent",
        "attempts_lvl1", "attempts_lvl2", "attempts_lvl3", "attempts_lvl4",
        "time_lvl1_s", "time_lvl2_s", "time_lvl3_s", "time_lvl4_s",
        "total_active_time_s", "total_hints_requested", "hints_effective_count",
        "tam_pu", "tam_peou", "tam_ai_scaffolding", "tam_ai_trust", "tam_intention_to_use", "sus_score"
    ]
    yield format_csv_row(headers)

    sessions_q = await db.execute(select(Session).order_by(Session.started_at.asc()))
    sessions = sessions_q.scalars().all()

    for s in sessions:
        # Attempts per level
        attempts = {}
        times = {}
        for lvl in range(1, 5):
            att_q = await db.execute(
                select(func.max(TelemetryEvent.attempt_number))
                .where(and_(TelemetryEvent.session_id == s.id, TelemetryEvent.level_id == lvl))
            )
            attempts[lvl] = att_q.scalar_one() or 0

            time_q = await db.execute(
                select(TelemetryEvent.active_time_ms)
                .where(and_(TelemetryEvent.session_id == s.id, TelemetryEvent.level_id == lvl, TelemetryEvent.event_type == "level_completed"))
            )
            t_ms = time_q.scalar_one()
            times[lvl] = round(t_ms / 1000, 2) if t_ms else 0.0

        # Total active time
        tot_time_q = await db.execute(
            select(func.sum(TelemetryEvent.active_time_ms))
            .where(and_(TelemetryEvent.session_id == s.id, TelemetryEvent.event_type == "level_completed"))
        )
        total_time_ms = tot_time_q.scalar_one() or 0

        # Hints
        hints_q = await db.execute(
            select(func.count(TelemetryEvent.id))
            .where(and_(TelemetryEvent.session_id == s.id, TelemetryEvent.event_type == "ai_hint_requested"))
        )
        total_hints = hints_q.scalar_one() or 0

        eff_q = await db.execute(
            select(func.count(TelemetryEvent.id))
            .where(and_(TelemetryEvent.session_id == s.id, TelemetryEvent.ai_hint_effective.is_(True)))
        )
        eff_hints = eff_q.scalar_one() or 0

        # Survey
        surv_q = await db.execute(select(SurveyResponse).where(SurveyResponse.session_id == s.id))
        surv = surv_q.scalar_one_or_none()

        row = [
            s.participant_id,
            s.condition,
            s.has_assent,
            attempts.get(1, 0),
            attempts.get(2, 0),
            attempts.get(3, 0),
            attempts.get(4, 0),
            times.get(1, 0.0),
            times.get(2, 0.0),
            times.get(3, 0.0),
            times.get(4, 0.0),
            round(total_time_ms / 1000, 2),
            total_hints,
            eff_hints,
            surv.tam_perceived_usefulness if surv else None,
            surv.tam_perceived_ease_of_use if surv else None,
            surv.tam_ai_scaffolding if surv else None,
            surv.tam_ai_trust if surv else None,
            surv.tam_intention_to_use if surv else None,
            surv.sus_score if surv else None,
        ]
        yield format_csv_row(row)

@router.get("/csv")
async def export_csv(
    type: str = Query("flat_events", description="Tipo de exportación: 'flat_events' o 'summary_by_student'"),
    db: AsyncSession = Depends(get_db),
):
    """Descarga de datasets en formato CSV plano."""
    if type == "summary_by_student":
        filename = "student_metrics_summary.csv"
        generator = stream_student_summary_csv(db)
    else:
        filename = "telemetry_events.csv"
        generator = stream_csv(db)

    return StreamingResponse(
        generator,
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )

@router.get("/jsonl")
async def export_jsonl(db: AsyncSession = Depends(get_db)):
    """Descarga masiva de eventos en formato JSON Lines."""
    return StreamingResponse(
        stream_jsonl(db),
        media_type="application/x-ndjson",
        headers={"Content-Disposition": "attachment; filename=telemetry_data.jsonl"},
    )


@router.get("/json")
async def export_json(db: AsyncSession = Depends(get_db)):
    """Descarga los eventos como un arreglo JSON estándar."""
    return StreamingResponse(
        stream_json(db),
        media_type="application/json",
        headers={"Content-Disposition": "attachment; filename=telemetry_data.json"},
    )

@router.get("/xlsx")
async def export_xlsx(db: AsyncSession = Depends(get_db)):
    """Descarga de dataset científico estructurado en Excel (4 hojas)."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="La librería 'openpyxl' no está disponible en este entorno.",
        )

    wb = openpyxl.Workbook()

    # Sheet 1: Participantes
    ws1 = wb.active
    ws1.title = "1_Participantes"
    ws1.append(["Participant ID", "Condición", "Asentimiento", "Inicio", "Fin"])

    sessions_q = await db.execute(select(Session).order_by(Session.started_at.asc()))
    sessions = sessions_q.scalars().all()
    for s in sessions:
        ws1.append([
            s.participant_id, s.condition, "Sí" if s.has_assent else "No",
            s.started_at.strftime("%Y-%m-%d %H:%M:%S") if s.started_at else "",
            s.completed_at.strftime("%Y-%m-%d %H:%M:%S") if s.completed_at else "",
        ])

    # Sheet 2: Rendimiento Niveles
    ws2 = wb.create_sheet(title="2_Rendimiento_Niveles")
    ws2.append(["Participant ID", "Nivel", "Intentos", "Tiempo Activo (s)", "Tiempo Inactivo (s)", "Éxito"])

    events_q = await db.execute(
        select(TelemetryEvent)
        .where(TelemetryEvent.event_type.in_(["level_completed", "code_run", "syntax_error"]))
        .order_by(TelemetryEvent.timestamp.asc())
    )
    events = events_q.scalars().all()
    for e in events:
        ws2.append([
            e.participant_id, e.level_id, e.attempt_number or 1,
            round((e.active_time_ms or 0) / 1000, 2),
            round((e.idle_time_ms or 0) / 1000, 2),
            "Sí" if e.is_success else "No",
        ])

    # Sheet 3: Uso IA Pistas
    ws3 = wb.create_sheet(title="3_Uso_IA_Pistas")
    ws3.append(["Participant ID", "Nivel", "Tipo de Pista", "Efectiva", "Fecha"])

    hints_q = await db.execute(
        select(TelemetryEvent)
        .where(TelemetryEvent.event_type == "ai_hint_requested")
        .order_by(TelemetryEvent.timestamp.asc())
    )
    hints = hints_q.scalars().all()
    for h in hints:
        ws3.append([
            h.participant_id, h.level_id, h.ai_hint_type or "N/A",
            "Sí" if h.ai_hint_effective else ("No" if h.ai_hint_effective is False else "Pendiente"),
            h.timestamp.strftime("%Y-%m-%d %H:%M:%S") if h.timestamp else "",
        ])

    # Sheet 4: Encuesta TAM SUS
    ws4 = wb.create_sheet(title="4_Encuesta_TAM_SUS")
    ws4.append(["Participant ID", "TAM Utilidad (PU)", "TAM Facilidad (PEOU)", "TAM Andamiaje IA", "TAM Confianza IA", "TAM Intención de Uso", "SUS Score (0-100)", "Fecha"])

    surv_q = await db.execute(select(SurveyResponse).order_by(SurveyResponse.submitted_at.asc()))
    surveys = surv_q.scalars().all()
    for surv in surveys:
        ws4.append([
            surv.participant_id, surv.tam_perceived_usefulness,
            surv.tam_perceived_ease_of_use, surv.tam_ai_scaffolding,
            surv.tam_ai_trust, surv.tam_intention_to_use,
            surv.sus_score if surv.sus_score is not None else "N/A",
            surv.submitted_at.strftime("%Y-%m-%d %H:%M:%S") if surv.submitted_at else "",
        ])

    # Save to memory buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=tivot_research_dataset.xlsx"},
    )
